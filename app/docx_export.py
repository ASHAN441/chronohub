"""按「值班表.docx」模板填充导出。

为什么是「填充」而不是「重新生成一份 docx」：
模板里有固定字号、列宽、5 条注意事项和一张图片。用代码重画必然走样，
最后拿到的是"看起来像但不一样"的东西。所以直接打开模板、只往单元格里写人名。

模板查找顺序（先用户可替换的，再打包内置的）：
1. exe 同级的 `值班表.docx` —— 想换模板不用重新打包
2. 打包内置的副本
"""

from __future__ import annotations

import os
import re

from .paths import base_dir, resource_dir

TEMPLATE_NAME = "值班表.docx"


def template_path() -> str:
    """返回可用的模板路径，找不到返回空串。"""
    candidates = [
        os.path.join(base_dir(), TEMPLATE_NAME),
        os.path.join(resource_dir(), TEMPLATE_NAME),
    ]
    for path in candidates:
        if os.path.exists(path):
            return path
    return ""


def _set_paragraph_text(paragraph, text: str) -> None:
    """替换段落文字，同时保留第一段 run 的格式（字号 / 字体）。"""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _set_cell_text(cell, text: str) -> None:
    """替换单元格文字，保留原有段落对齐与首个 run 的格式。"""
    paragraph = cell.paragraphs[0]
    _set_paragraph_text(paragraph, text)
    # 单元格里若被塞了多个段落，清掉多余的
    for extra in cell.paragraphs[1:]:
        _set_paragraph_text(extra, "")


def _row_label(row) -> str:
    """第一列的文字，去掉空白便于匹配。"""
    try:
        text = row.cells[0].text or ""
    except IndexError:
        return ""
    return re.sub(r"\s+", "", text)


def _block_start(label: str):
    """从「第1、2节」里取出起始节次；双岗行返回 None。"""
    if "台务" in label or "部长" in label:
        return None
    match = re.search(r"第(\d+)", label)
    if match:
        return int(match.group(1))
    return None


def export_roster(roster: dict, names: dict, out_path: str, template: str = None) -> str:
    """把一张值班表填进模板并另存。

    roster  {"week": 3, "assignments": [...]}
    names   {人员key: 显示名}
    """
    from docx import Document

    source = template or template_path()
    if not source:
        raise FileNotFoundError(
            f"找不到值班表模板「{TEMPLATE_NAME}」，请把它放在程序所在目录"
        )

    document = Document(source)
    week = roster.get("week")

    # 1) 标题里的周次
    for paragraph in document.paragraphs:
        text = paragraph.text or ""
        if "值班记录" in text and "第" in text:
            new_text = re.sub(r"第\s*[（(]\s*[）)]\s*周", f"第（{week}）周", text)
            if new_text == text:
                new_text = re.sub(r"第[（(][^）)]*[）)]周", f"第（{week}）周", text)
            _set_paragraph_text(paragraph, new_text)
            break

    # 2) 表格
    if not document.tables:
        raise ValueError("模板里没有表格，可能不是预期的值班表模板")

    table = document.tables[0]

    # 按「星期 + 节次」索引本次排班结果
    by_day_block = {}
    by_day_dual = {}
    for item in roster.get("assignments") or []:
        if item.get("kind") == "dual":
            by_day_dual[item["day"]] = item
        else:
            by_day_block[(item["day"], item.get("start"))] = item

    def name_of(key: str) -> str:
        key = str(key or "").strip()
        if not key:
            return ""
        return names.get(key, key)

    filled = 0
    for row in table.rows:
        label = _row_label(row)
        if not label:
            continue
        # 跳过「联系方式」「值班记录」等非人名行
        if "联系方式" in label or "值班记录" in label:
            continue
        if "星期" in label:
            continue

        cells = row.cells
        if len(cells) < 6:
            continue

        start = _block_start(label)
        is_dual = start is None

        for day in range(1, 6):
            if day >= len(cells):
                break
            cell = cells[day]
            if is_dual:
                item = by_day_dual.get(day)
                hall = name_of(item.get("hall")) if item else ""
                minister = name_of(item.get("minister")) if item else ""
                text = ""
                if hall or minister:
                    text = f"台务：{hall}　部长：{minister}"
                _set_cell_text(cell, text)
                if text:
                    filled += 1
            else:
                item = by_day_block.get((day, start))
                text = name_of(item.get("person")) if item else ""
                _set_cell_text(cell, text)
                if text:
                    filled += 1

    directory = os.path.dirname(os.path.abspath(out_path))
    if directory:
        os.makedirs(directory, exist_ok=True)

    document.save(out_path)
    return out_path
